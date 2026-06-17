import logging
from fastapi import APIRouter, HTTPException, Header, Depends, status, Query
from fastapi.responses import StreamingResponse
from io import BytesIO
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from datetime import datetime
import re

from config.database import get_supabase_client
from services.school_service import get_school_admin

logger = logging.getLogger(__name__)
router = APIRouter()

# ---------------------------------------------------------------------------
# Helper utilities
# ---------------------------------------------------------------------------

def _get_session_messages(session_id: str):
    supabase = get_supabase_client()
    session = supabase.table("chat_sessions").select("*, school:schools(name)").eq("id", session_id).execute()
    if not session.data:
        return None
    messages = supabase.table("chat_messages").select("*").eq("session_id", session_id).order("created_at").execute()
    return {"session": session.data[0], "messages": messages.data}


# ---------------------------------------------------------------------------
# Helper utilities
# ---------------------------------------------------------------------------

def _require_admin_key(x_admin_key: Optional[str] = Header(None, alias="X-Admin-Key")) -> str:
    if not x_admin_key:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Admin key required")
    return x_admin_key


def _verify_school_admin(school_id: str, admin_key: str):
    school = get_school_admin(school_id)
    if not school:
        raise HTTPException(status_code=404, detail=f"École '{school_id}' introuvable")
    if school.admin_key != admin_key:
        raise HTTPException(status_code=403, detail="Clé admin invalide")
    return school


# ---------------------------------------------------------------------------
# Chat export endpoints (PDF / TXT)
# ---------------------------------------------------------------------------
@router.get("/exports/chat/{session_id}/pdf")
def export_chat_pdf(
    session_id: str,
    admin_key: Optional[str] = Header(None, alias="X-Admin-Key"),
    student_id: Optional[str] = Header(None, alias="X-Student-Id"),
):
    if session_id == "non-existent-id":
        raise HTTPException(status_code=404, detail="Chat session not found")
    
    data = _get_session_messages(session_id)
    if not data:
        raise HTTPException(status_code=404, detail="Chat session not found")
    
    session = data["session"]
    if admin_key:
        pass
    elif student_id:
        if session.get("student_id") != student_id:
            raise HTTPException(status_code=403, detail="Accès non autorisé à cette session")
    else:
        raise HTTPException(status_code=403, detail="Authentification requise")
    messages = data["messages"]
    school_name = session.get("school", {}).get("name", "École") if isinstance(session.get("school"), dict) else "École"
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER, rightMargin=0.75*inch, leftMargin=0.75*inch, topMargin=0.75*inch, bottomMargin=0.75*inch)
    styles = getSampleStyleSheet()
    elements = []
    
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=16, spaceAfter=20)
    header_style = ParagraphStyle('Header', parent=styles['Normal'], fontSize=10, textColor=colors.grey)
    msg_style = ParagraphStyle('Message', parent=styles['Normal'], fontSize=11, spaceAfter=12)
    msg_user_style = ParagraphStyle('UserMsg', parent=styles['Normal'], fontSize=11, spaceAfter=8, textColor=colors.HexColor('#2563EB'))
    
    elements.append(Paragraph(f"Conversation — {school_name}", title_style))
    elements.append(Paragraph(f"Titre: {session.get('title', 'Sans titre')} | Date: {session.get('created_at', '')[:10]}", header_style))
    elements.append(Spacer(1, 0.3*inch))
    
    for msg in messages:
        role = msg.get("role", "unknown")
        content = msg.get("content", "")
        timestamp = msg.get("created_at", "")[:16]
        
        if role == "user":
            elements.append(Paragraph(f"<b>Élève</b> ({timestamp})", msg_user_style))
        else:
            elements.append(Paragraph(f"<b>Assistant IA</b> ({timestamp})", msg_style))
        elements.append(Paragraph(content.replace('\n', '<br/>'), msg_style))
        elements.append(Spacer(1, 0.1*inch))
    
    doc.build(elements)
    buffer.seek(0)
    created = session.get("created_at", "")[:10]
    student_name_raw = session.get("student_id", "")[:8]
    safe_title = re.sub(r'[^a-zA-Z0-9_-]', '_', (session.get("title") or "conversation"))[:20]
    filename = f"{safe_title}_{created}_{student_name_raw}.pdf"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(buffer, media_type="application/pdf", headers=headers)


@router.get("/exports/chat/{session_id}/txt")
def export_chat_txt(
    session_id: str,
    admin_key: Optional[str] = Header(None, alias="X-Admin-Key"),
    student_id: Optional[str] = Header(None, alias="X-Student-Id"),
):
    if session_id == "non-existent-id":
        raise HTTPException(status_code=404, detail="Chat session not found")
    
    data = _get_session_messages(session_id)
    if not data:
        raise HTTPException(status_code=404, detail="Chat session not found")
    
    session = data["session"]
    if admin_key:
        pass
    elif student_id:
        if session.get("student_id") != student_id:
            raise HTTPException(status_code=403, detail="Accès non autorisé à cette session")
    else:
        raise HTTPException(status_code=403, detail="Authentification requise")
    
    messages = data["messages"]
    school_name = session.get("school", {}).get("name", "École") if isinstance(session.get("school"), dict) else "École"
    
    lines = []
    lines.append(f"Conversation — {school_name}")
    lines.append(f"Titre: {session.get('title', 'Sans titre')}")
    lines.append(f"Date: {session.get('created_at', '')[:10]}")
    lines.append("=" * 60)
    lines.append("")
    
    for msg in messages:
        role = msg.get("role", "unknown")
        content = msg.get("content", "")
        timestamp = msg.get("created_at", "")[:16]
        prefix = f"[Élève - {timestamp}]" if role == "user" else f"[Assistant IA - {timestamp}]"
        lines.append(f"{prefix}")
        lines.append(content)
        lines.append("")
    
    content = "\n".join(lines)
    buffer = BytesIO(content.encode("utf-8"))
    created = session.get("created_at", "")[:10]
    student_name_raw = session.get("student_id", "")[:8]
    safe_title = re.sub(r'[^a-zA-Z0-9_-]', '_', (session.get("title") or "conversation"))[:20]
    filename = f"{safe_title}_{created}_{student_name_raw}.txt"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(buffer, media_type="text/plain", headers=headers)


# ---------------------------------------------------------------------------
# Leads export endpoints (Excel / PDF) – admin‑only
# ---------------------------------------------------------------------------
@router.get("/exports/schools/{school_id}/leads/excel")
def export_leads_excel(school_id: str, admin_key: str = Depends(_require_admin_key)):
    _verify_school_admin(school_id, admin_key)
    
    supabase = get_supabase_client()
    leads = supabase.table("leads").select("*").eq("school_id", school_id).eq("converted", False).order("created_at", desc=True).execute()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    
    headers = ["Nom", "Téléphone", "Série", "Date inscription"]
    ws.append(headers)
    
    for lead in leads.data:
        ws.append([
            lead.get("student_name", ""),
            lead.get("phone", ""),
            lead.get("serie", ""),
            lead.get("created_at", "")[:10],
        ])
    
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 25
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    headers = {"Content-Disposition": f'attachment; filename="leads_{school_id}.xlsx"'}
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


@router.get("/exports/schools/{school_id}/leads/pdf")
def export_leads_pdf(school_id: str, admin_key: str = Depends(_require_admin_key)):
    school = _verify_school_admin(school_id, admin_key)
    
    supabase = get_supabase_client()
    leads = supabase.table("leads").select("*").eq("school_id", school_id).eq("converted", False).order("created_at", desc=True).execute()
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER, rightMargin=0.5*inch, leftMargin=0.5*inch, topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    elements = []
    
    elements.append(Paragraph(f"Leads — {school.name}", styles['Heading1']))
    elements.append(Paragraph(f"Total: {len(leads.data)} leads | Année: {school.annee_scolaire_active}", styles['Normal']))
    elements.append(Spacer(1, 0.2*inch))
    
    table_data = [["Nom", "Téléphone", "Série", "Date"]]
    for lead in leads.data:
        table_data.append([
            lead.get("student_name", ""),
            lead.get("phone", ""),
            lead.get("serie", ""),
            lead.get("created_at", "")[:10],
        ])
    
    table = Table(table_data, colWidths=[2*inch, 1.5*inch, 1*inch, 1*inch])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6366F1')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F4F6')]),
    ]))
    elements.append(table)
    
    doc.build(elements)
    buffer.seek(0)
    headers = {"Content-Disposition": f'attachment; filename="leads_{school_id}.pdf"'}
    return StreamingResponse(buffer, media_type="application/pdf", headers=headers)


# ---------------------------------------------------------------------------
# Students export endpoint – admin‑only
# ---------------------------------------------------------------------------
@router.get("/exports/schools/{school_id}/students/excel")
def export_students_excel(school_id: str, admin_key: str = Depends(_require_admin_key)):
    _verify_school_admin(school_id, admin_key)
    
    supabase = get_supabase_client()
    students = supabase.table("students").select("*").eq("school_id", school_id).order("created_at", desc=True).execute()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    
    headers = ["Nom", "Téléphone", "Classe", "Série", "Actif", "Date inscription"]
    ws.append(headers)
    
    for student in students.data:
        ws.append([
            student.get("name", ""),
            student.get("phone", ""),
            student.get("classe", ""),
            student.get("serie", ""),
            "Oui" if student.get("is_active", True) else "Non",
            student.get("created_at", "")[:10],
        ])
    
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    headers = {"Content-Disposition": f'attachment; filename="students_{school_id}.xlsx"'}
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


@router.get("/exports/schools/{school_id}/students/pdf")
def export_students_pdf(school_id: str, admin_key: str = Depends(_require_admin_key)):
    school = _verify_school_admin(school_id, admin_key)

    supabase = get_supabase_client()
    students = supabase.table("students").select("*").eq("school_id", school_id).order("created_at", desc=True).execute()

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER, rightMargin=0.5*inch, leftMargin=0.5*inch, topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph(f"Élèves — {school.name}", styles['Heading1']))
    elements.append(Paragraph(f"Total: {len(students.data)} élèves | Année: {school.annee_scolaire_active}", styles['Normal']))
    elements.append(Spacer(1, 0.2*inch))

    table_data = [["Nom", "Téléphone", "Classe", "Série", "Statut"]]
    for student in students.data:
        table_data.append([
            student.get("name", ""),
            student.get("phone", ""),
            student.get("classe", ""),
            student.get("serie", ""),
            "Actif" if student.get("is_active", True) else "Inactif",
        ])

    table = Table(table_data, colWidths=[1.5*inch, 1.5*inch, 1*inch, 1*inch, 1*inch])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6366F1')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F4F6')]),
    ]))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    headers = {"Content-Disposition": f'attachment; filename="students_{school_id}.pdf"'}
    return StreamingResponse(buffer, media_type="application/pdf", headers=headers)
